import csv, json
from .engine import DISCLAIMER

def _cell(v):
    if isinstance(v, (list, tuple)):
        return ";".join(f"{x:.1f}" if isinstance(x, float) else str(x) for x in v)
    return v


CSV_COLS = ["id", "rule", "category", "severity", "description", "expected", "detected", "difference", "identifier", "layout",
            "floor_level", "schedule_ref", "entity_handles", "bbox", "confidence", "reason", "recommended_action", "review_status", "sub_check"]

def to_json(result, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(result.to_dict(), f, indent=2, default=str)

def to_csv(result, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["# " + DISCLAIMER])
        w.writerow(CSV_COLS)
        for fd in result.findings:
            d = fd.to_dict()
            w.writerow([_cell(d[c]) for c in CSV_COLS])

def to_xlsx(result, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active; ws.title = "Summary"
    orange = PatternFill("solid", fgColor="F58220")
    ws.append(["StructCheck report"]); ws["A1"].font = Font(bold=True, size=14, color="D96B10")
    for k, v in result.drawing.items():
        if not isinstance(v, (dict, list)): ws.append([k, v])
    ws.append([]); ws.append(["Source SHA-256 before", result.integrity["sha256_before"]])
    ws.append(["Source SHA-256 after", result.integrity["sha256_after"]]); ws.append(["Source unchanged", result.integrity["unchanged"]])
    ws.append([]); [ws.append([k, v]) for k, v in result.summary().items()]
    ws.append([]); ws.append(["Rule", "Name", "Status", "Findings", "Not-checked items", "Notes"])
    for c in ws[ws.max_row]: c.font = Font(bold=True, color="FFFFFF"); c.fill = orange
    for r in result.rules: ws.append([r["rule"], r["name"], r["status"], r["finding_count"], r["not_checked_count"], " | ".join(r["notes"])])
    ws.append([]); ws.append([DISCLAIMER])
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 60
    wf = wb.create_sheet("Findings"); wf.append(CSV_COLS)
    for c in wf[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = orange
    colour = {"ERROR": "FDE2E0", "WARNING": "FDE8D3", "REVIEW": "FBF0CF", "NOT CHECKED": "E5E7EB"}
    for fd in result.findings:
        d = fd.to_dict(); wf.append([_cell(d[c]) for c in CSV_COLS])
        wf.cell(wf.max_row, 4).fill = PatternFill("solid", fgColor=colour.get(d["severity"], "FFFFFF"))
    wf.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJ", (10, 6, 16, 12, 60, 18, 18, 14, 16, 12)): wf.column_dimensions[col].width = w
    wt = wb.create_sheet("Text Height Analysis")
    rows = result.text_height_analysis.get("analysis", [])
    cols = ["layout", "category", "text", "layer", "style", "kind", "handle", "raw_height", "effective_height", "effective_certain",
            "expected", "expected_source", "difference", "included", "exclusion_reason", "severity", "review_status"]
    wt.append(cols)
    for c in wt[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = orange
    for r in rows: wt.append([_cell(r.get(c)) for c in cols])
    wt.append([]); wt.append(["Distribution of effective heights by category"])
    for k, v in result.text_height_analysis.get("distribution", {}).items(): wt.append([k, str(v)])
    wb.save(path)
